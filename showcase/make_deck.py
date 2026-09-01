"""Generate the final-showcase PowerPoint deck (Ticket 13).

Reproducibly builds ``showcase/AgentTrail-final-showcase.pptx`` — exactly
three high-contrast 16:9 slides for the strict two-minute showcase,
specified word for word in ``showcase/deck-content.md``. Every figure is
taken from verified command output on this branch; nothing is invented.

This is a presentation-authoring tool, NOT part of the AgentTrail
application. It requires ``python-pptx``, which must stay out of the
application's dependency set (the offline application paths are
deliberately standard-library-only). Run it from an isolated virtual
environment:

    python3 -m venv /tmp/agenttrail-deckenv
    /tmp/agenttrail-deckenv/bin/pip install python-pptx
    /tmp/agenttrail-deckenv/bin/python showcase/make_deck.py

The script writes the deck next to itself and prints the output path.
"""

from pathlib import Path

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_CONNECTOR, MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.oxml.ns import qn
from pptx.util import Inches, Pt

# Palette: near-black ground, white type, one electric-blue accent.
# Safety green appears only on the v1 result; strong red only on the
# blocked v2 result. Nothing decorative.
BG = RGBColor(0x0A, 0x0E, 0x14)  # near-black
WHITE = RGBColor(0xFF, 0xFF, 0xFF)
MUTED = RGBColor(0x93, 0xA1, 0xAD)  # supporting text on dark
FAINT = RGBColor(0x62, 0x6E, 0x79)  # footers, small print
BLUE = RGBColor(0x3D, 0xA5, 0xFF)  # electric blue — audit trail / regulation
PANEL = RGBColor(0x16, 0x1D, 0x27)  # flow boxes
PANEL_LINE = RGBColor(0x2E, 0x3A, 0x47)
GREEN = RGBColor(0x2E, 0xC2, 0x6E)  # v1 PASS only
RED = RGBColor(0xF1, 0x4C, 0x4C)  # v2 BLOCKED only

FONT = "Helvetica Neue"

SLIDE_W = Inches(13.333)
SLIDE_H = Inches(7.5)


def _fill_frame(
    frame,
    lines,
    size=18,
    color=WHITE,
    bold=False,
    align=PP_ALIGN.LEFT,
    anchor=MSO_ANCHOR.TOP,
    line_gap=None,
):
    """Fill a text frame; each line is text or (text, overrides).

    ``text`` may itself be a list of (text, overrides) runs so a single
    paragraph can mix colours (used for the accented headline phrase).
    """
    frame.word_wrap = True
    frame.vertical_anchor = anchor
    for i, entry in enumerate(lines):
        text, overrides = entry if isinstance(entry, tuple) else (entry, {})
        para = frame.paragraphs[0] if i == 0 else frame.add_paragraph()
        para.alignment = overrides.get("align", align)
        if line_gap is not None and i > 0:
            para.space_before = Pt(line_gap)
        runs = text if isinstance(text, list) else [(text, {})]
        for run_text, run_overrides in runs:
            merged = {**overrides, **run_overrides}
            run = para.add_run()
            run.text = run_text
            font = run.font
            font.name = FONT
            font.size = Pt(merged.get("size", size))
            font.bold = merged.get("bold", bold)
            font.italic = merged.get("italic", False)
            font.color.rgb = merged.get("color", color)


def _text_box(slide, x, y, w, h, lines, **kwargs):
    box = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    _fill_frame(box.text_frame, lines, **kwargs)
    return box


def _shape_box(
    slide,
    x,
    y,
    w,
    h,
    lines,
    fill=PANEL,
    text_color=WHITE,
    border=PANEL_LINE,
    size=16,
    bold=False,
    align=PP_ALIGN.CENTER,
):
    shape = slide.shapes.add_shape(
        MSO_SHAPE.ROUNDED_RECTANGLE, Inches(x), Inches(y), Inches(w), Inches(h)
    )
    shape.adjustments[0] = 0.14
    if fill is None:
        shape.fill.background()
    else:
        shape.fill.solid()
        shape.fill.fore_color.rgb = fill
    if border is None:
        shape.line.fill.background()
    else:
        shape.line.color.rgb = border
        shape.line.width = Pt(1.25)
    shape.shadow.inherit = False
    _fill_frame(
        shape.text_frame,
        lines,
        size=size,
        color=text_color,
        bold=bold,
        align=align,
        anchor=MSO_ANCHOR.MIDDLE,
    )
    return shape


def _arrow(slide, x1, y1, x2, y2, color=MUTED, weight=1.75):
    conn = slide.shapes.add_connector(
        MSO_CONNECTOR.STRAIGHT, Inches(x1), Inches(y1), Inches(x2), Inches(y2)
    )
    conn.line.color.rgb = color
    conn.line.width = Pt(weight)
    ln = conn.line._get_or_add_ln()
    tail = ln.makeelement(
        qn("a:tailEnd"), {"type": "triangle", "w": "med", "len": "med"}
    )
    ln.append(tail)
    return conn


def _blank_slide(prs):
    slide = prs.slides.add_slide(prs.slide_layouts[6])  # blank layout
    fill = slide.background.fill
    fill.solid()
    fill.fore_color.rgb = BG
    return slide


def _brand_label(slide):
    _text_box(
        slide,
        0.9,
        0.45,
        4.0,
        0.4,
        [("AGENTTRAIL", {"size": 14, "bold": True, "color": BLUE})],
    )


def _notes(slide, text):
    slide.notes_slide.notes_text_frame.text = text


def slide_1_hook(prs):
    slide = _blank_slide(prs)
    _brand_label(slide)

    _text_box(
        slide,
        0.9,
        1.55,
        11.6,
        3.2,
        [
            ("By December 2027,", {}),
            ([("‘the AI said so’", {"color": BLUE})], {}),
            ("will no longer be enough.", {}),
        ],
        size=54,
        bold=True,
        color=WHITE,
        line_gap=4,
    )

    _text_box(
        slide,
        0.9,
        4.95,
        10.8,
        0.75,
        [
            (
                "For many high-risk uses, Article 12 requires AI systems "
                "to support automatic event logging and traceability.",
                {"size": 19, "color": MUTED},
            )
        ],
    )

    # Single motif: a trail of audit events — a thin blue baseline with
    # event markers, the newest one lit.
    trail_y = 6.05
    slide.shapes.add_shape(
        MSO_SHAPE.RECTANGLE,
        Inches(0.9),
        Inches(trail_y),
        Inches(11.53),
        Pt(1.5),
    )
    baseline = slide.shapes[-1]
    baseline.fill.solid()
    baseline.fill.fore_color.rgb = PANEL_LINE
    baseline.line.fill.background()
    baseline.shadow.inherit = False
    marker_xs = [1.6, 3.55, 5.5, 7.45, 9.4, 11.35]
    for i, mx in enumerate(marker_xs):
        last = i == len(marker_xs) - 1
        size = 0.22 if last else 0.14
        marker = slide.shapes.add_shape(
            MSO_SHAPE.OVAL,
            Inches(mx - size / 2),
            Inches(trail_y + 0.01 - size / 2),
            Inches(size),
            Inches(size),
        )
        marker.fill.solid()
        marker.fill.fore_color.rgb = BLUE if last else PANEL_LINE
        if last:
            marker.line.fill.background()
        else:
            marker.line.color.rgb = RGBColor(0x3E, 0x4C, 0x5A)
            marker.line.width = Pt(0.75)
        marker.shadow.inherit = False

    _text_box(
        slide,
        0.9,
        6.85,
        11.6,
        0.4,
        [
            (
                "EU AI Act  ·  Annex III high-risk rules apply "
                "2 December 2027  ·  engineering alignment, "
                "not a compliance claim",
                {"size": 11, "color": FAINT},
            )
        ],
    )

    _notes(
        slide,
        "SPEAKER SCRIPT — READ THIS\n"
        "\n"
        "From December 2027, for many high-risk uses, “the AI said "
        "so” will not be enough.\n"
        "\n"
        "Article 12 of the EU AI Act requires systems to support "
        "automatic logging, so important decisions can be traced.\n"
        "\n"
        "AgentTrail turns that principle into an engineering "
        "control.\n"
        "\n"
        "PRESENTER CUE: Advance to Slide 2.\n"
        "\n"
        "SOURCES AND ACCURACY NOTES\n"
        "\n"
        "- EU AI Act (Reg. (EU) 2024/1689) generally applicable "
        "2 Aug 2026.\n"
        "- Per Reg. (EU) 2026/1744, Annex III high-risk rules apply "
        "from 2 Dec 2027.\n"
        "- High-risk rules for Annex I embedded products apply from "
        "2 Aug 2028.\n"
        "- Art. 12: high-risk AI systems must technically allow "
        "automatic recording of events over their lifetime, "
        "supporting traceability appropriate to the system's "
        "purpose.\n"
        "- AgentTrail = engineering alignment, not a legal-compliance "
        "claim; this demonstrator is not a regulated high-risk "
        "system.\n"
        "- Sources: digital-strategy.ec.europa.eu/en/policies/"
        "regulatory-framework-ai · eur-lex.europa.eu/eli/reg/2026/"
        "1744/oj/eng · eur-lex.europa.eu/eli/reg/2024/1689/oj/eng",
    )


def slide_2_architecture(prs):
    slide = _blank_slide(prs)

    _text_box(
        slide,
        0.9,
        0.35,
        11.6,
        1.5,
        [
            (
                "AgentTrail turns evidence into a release decision.",
                {"size": 44, "bold": True, "color": WHITE},
            )
        ],
    )

    chain_x, chain_w = 4.5, 4.35
    centre = chain_x + chain_w / 2
    gap = 0.26

    y = 2.0
    _shape_box(slide, chain_x, y, chain_w, 0.55, [("AI decision", {})])
    y += 0.55
    _arrow(slide, centre, y, centre, y + gap)
    y += gap

    trail_y, trail_h = y, 0.8
    _shape_box(
        slide,
        chain_x,
        trail_y,
        chain_w,
        trail_h,
        [("append-only audit trail", {"size": 19, "bold": True})],
        fill=BLUE,
        text_color=BG,
        border=None,
    )
    y += trail_h
    _arrow(slide, centre, y, centre, y + gap)
    y += gap

    eval_y, eval_h = y, 0.55
    _shape_box(
        slide, chain_x, eval_y, chain_w, eval_h, [("behavioural evaluation", {})]
    )
    y += eval_h
    _arrow(slide, centre, y, centre, y + gap)
    y += gap

    _shape_box(
        slide,
        chain_x,
        y,
        chain_w,
        0.55,
        [("CI safety gate", {"bold": True})],
        fill=WHITE,
        text_color=BG,
        border=None,
    )

    # Labelled cases and recorded responses feed the evaluation.
    _shape_box(
        slide,
        0.9,
        eval_y - 0.14,
        2.9,
        0.88,
        [
            ("labelled golden dataset", {"size": 13, "color": MUTED}),
            ("recorded responses", {"size": 13, "color": MUTED}),
        ],
        fill=None,
        border=PANEL_LINE,
    )
    _arrow(
        slide,
        3.8,
        eval_y + eval_h / 2,
        chain_x - 0.06,
        eval_y + eval_h / 2,
    )

    # Policy replay branches from the trail — read-only.
    _shape_box(
        slide,
        10.0,
        trail_y + 0.11,
        2.5,
        0.62,
        [("policy replay", {"size": 15, "bold": True, "color": BLUE})],
        fill=None,
        border=BLUE,
    )
    _arrow(
        slide,
        chain_x + chain_w,
        trail_y + trail_h / 2,
        9.94,
        trail_y + trail_h / 2,
        color=BLUE,
    )
    _text_box(
        slide,
        10.0,
        trail_y + 0.85,
        2.7,
        0.4,
        [("reads the same trail", {"size": 11, "color": FAINT})],
    )

    _text_box(
        slide,
        0.9,
        5.95,
        11.6,
        0.55,
        [
            (
                "Offline. Deterministic. No live model call in CI.",
                {"size": 18, "bold": True, "color": WHITE},
            )
        ],
        align=PP_ALIGN.CENTER,
    )
    _text_box(
        slide,
        0.9,
        6.85,
        11.6,
        0.4,
        [
            (
                "trail is a local append-only JSONL file — "
                "a demonstrator, not tamper-proof production storage",
                {"size": 11, "color": FAINT},
            )
        ],
        align=PP_ALIGN.CENTER,
    )

    _notes(
        slide,
        "SPEAKER SCRIPT — READ THIS\n"
        "\n"
        "Here’s how.\n"
        "\n"
        "The agent makes a decision, and AgentTrail stores it in an "
        "append-only audit trail.\n"
        "\n"
        "The evaluation reads that decision back from the trail, "
        "compares it with the expected behaviour, and the CI gate "
        "decides whether the change can be released.\n"
        "\n"
        "Recorded model responses keep the test deterministic and "
        "offline.\n"
        "\n"
        "PRESENTER CUE: Switch to Terminal.\n"
        "\n"
        "First, the safe version.\n"
        "\n"
        "Run:\n"
        "\n"
        "python3 -m agenttrail.ci_gate --agent-version v1\n"
        "\n"
        "It escalates all twelve low-confidence cases, so the gate "
        "passes.\n"
        "\n"
        "Now I’ll run an intentionally unsafe version that has been "
        "made over-confident.\n"
        "\n"
        "Run:\n"
        "\n"
        "python3 -m agenttrail.ci_gate --agent-version v2\n"
        "\n"
        "PRESENTER CUE: Stop typing and pause when the red result "
        "appears.\n"
        "\n"
        "It escalates only nine cases and misses three decisions that "
        "should have gone to a human.\n"
        "\n"
        "The process exits with status one and blocks the pipeline.\n"
        "\n"
        "This red result is not the demo failing. It is the safety "
        "control doing its job: stopping an unsafe change before "
        "deployment.\n"
        "\n"
        "PRESENTER CUE: Return to Slide 3.\n"
        "\n"
        "SOURCES AND ACCURACY NOTES\n"
        "\n"
        "- Evaluation reads decisions back from the persisted "
        "append-only audit trail, never the agent's return value.\n"
        "- Runs are deterministic and offline against recorded "
        "responses — no live model call in CI.\n"
        "- Expected live outputs: v1 → CI GATE: PASS, "
        "measured_rate=1.0000, 12/12 escalated, exit 0. v2 → "
        "CI GATE: FAIL, measured_rate=0.7500, 9/12, three mismatch "
        "lines, Decision: Gate blocked, exit 1 (intentional — v2 is "
        "a synthetic, controlled unsafe version).\n"
        "- The local JSONL trail is a demonstrator, not tamper-proof "
        "production storage.",
    )


def slide_3_result(prs):
    slide = _blank_slide(prs)

    _text_box(
        slide,
        0.9,
        0.35,
        11.6,
        1.5,
        [
            (
                "Unsafe change: blocked before deployment.",
                {"size": 44, "bold": True, "color": WHITE},
            )
        ],
    )

    rows = [
        (2.0, "v1", "100%", "PASS", GREEN),
        (3.35, "v2", "75%", "BLOCKED", RED),
    ]
    for row_y, version, rate, verdict, color in rows:
        _text_box(
            slide,
            1.1,
            row_y + 0.25,
            1.4,
            0.8,
            [(version, {"size": 30, "bold": True, "color": MUTED})],
        )
        _text_box(
            slide,
            2.7,
            row_y,
            4.6,
            1.3,
            [(rate, {"size": 80, "bold": True, "color": color})],
        )
        _shape_box(
            slide,
            8.4,
            row_y + 0.26,
            3.4,
            0.78,
            [(verdict, {"size": 26, "bold": True})],
            fill=color,
            text_color=BG,
            border=None,
        )

    _text_box(
        slide,
        0.9,
        4.85,
        11.6,
        0.45,
        [
            (
                "three missed escalations  ·  exit status 1 stops "
                "the pipeline",
                {"size": 17, "color": MUTED},
            )
        ],
        align=PP_ALIGN.CENTER,
    )

    _text_box(
        slide,
        0.9,
        5.45,
        11.6,
        0.6,
        [
            (
                "36 past decisions replayed. One would change.",
                {"size": 28, "bold": True, "color": BLUE},
            )
        ],
        align=PP_ALIGN.CENTER,
    )

    _text_box(
        slide,
        0.9,
        6.3,
        11.6,
        0.85,
        [
            (
                "An unsafe AI change cannot pass silently.",
                {"size": 34, "bold": True, "color": WHITE},
            )
        ],
        align=PP_ALIGN.CENTER,
    )

    _notes(
        slide,
        "SPEAKER SCRIPT — READ THIS\n"
        "\n"
        "And because the trail exists, I can test a proposed policy "
        "against past decisions before applying it.\n"
        "\n"
        "Across thirty-six records, one decision would change.\n"
        "\n"
        "That gives people evidence for review. It does not "
        "automatically approve the policy.\n"
        "\n"
        "The important result is simple:\n"
        "\n"
        "An unsafe AI change cannot pass silently.\n"
        "\n"
        "PRESENTER CUE: Stop there. No thank-you slide; keep "
        "future work for questions.\n"
        "\n"
        "SOURCES AND ACCURACY NOTES\n"
        "\n"
        "- All figures from verified runs on this branch: v1 gate "
        "PASS (measured_rate=1.0000, 12/12 escalated, exit 0); v2 "
        "gate FAIL (measured_rate=0.7500, 9/12, mismatches SYN-003 / "
        "SYN-015 / SYN-027, exit 1 — intentional controlled "
        "regression).\n"
        "- Policy replay: 36 records evaluated, 35 unchanged, 1 flip "
        "— SYN-031 decide → escalate under the proposed 0.95 "
        "confidence threshold (stored confidence 0.9000).\n"
        "- Replay is read-only change-impact evidence, not approval "
        "or rejection of the policy.",
    )


def build(output_path):
    prs = Presentation()
    prs.slide_width = SLIDE_W
    prs.slide_height = SLIDE_H
    slide_1_hook(prs)
    slide_2_architecture(prs)
    slide_3_result(prs)
    prs.save(output_path)
    return output_path


if __name__ == "__main__":
    out = Path(__file__).resolve().parent / "AgentTrail-final-showcase.pptx"
    print(build(str(out)))
